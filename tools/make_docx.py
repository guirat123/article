#!/usr/bin/env python3
"""Export paper.tex to an editable Word document (build/manuscript_editable.docx).

The LaTeX source remains authoritative; this export exists so co-authors can
comment and edit in Word.  Mathematics is flattened to readable plain text.
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def flatten(tex):
    s = tex
    s = re.sub(r"\\todo\{([^}]*)\}", r"[TODO: \1]", s)
    s = re.sub(r"\\cite\{([^}]*)\}", r"[\1]", s)
    s = re.sub(r"\\ref\{([^}]*)\}", r"<\1>", s)
    s = s.replace("\\label", "")
    s = re.sub(r"\\(textbf|emph|textit|texttt)\{([^}]*)\}", r"\2", s)
    s = re.sub(r"\\(inst|orcidID|email)\{([^}]*)\}", r"\2", s)
    s = s.replace("\\and", " | ")
    s = re.sub(r"\$([^$]*)\$", r" \1 ", s)          # inline math -> text
    for a, b in [(r"\\hat", "^"), (r"\\mathbf", ""), (r"\\mathrm", ""),
                 (r"\\mathcal", ""), (r"\\sum", "SUM"), (r"\\frac", ""),
                 (r"\\lambda", "lambda"), (r"\\mu", "mu"), (r"\\sigma", "sigma"),
                 (r"\\xi", "xi"), (r"\\phi", "phi"), (r"\\Omega", "Omega"),
                 (r"\\in", "in"), (r"\\geq", ">="), (r"\\leq", "<="),
                 (r"\\rightarrow", "->"), (r"\\approx", "~"), (r"\\dagger", "+"),
                 (r"\\times", "x"), (r"\\approx", "~"), (r"\{", ""), (r"\}", ""),
                 ("_", ""), ("^", "")]:
        s = s.replace(a, b)
    s = re.sub(r"\\[,;! ]", " ", s)
    s = s.replace("~", " ").replace("---", "-").replace("--", "-")
    s = s.replace("``", '"').replace("''", '"')
    s = s.replace("\\%", "%").replace("\\&", "&").replace("\\#", "#")
    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = re.sub(r"[{}]", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def main():
    src = open(os.path.join(ROOT, "paper.tex")).read()
    src = re.sub(r"(?<!\\)%[^\n]*", "", src)
    body = src[src.find("\\begin{document}"):src.rfind("\\end{document}")]

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    for name, size, bold in [("Heading 1", 14, True), ("Heading 2", 12, True),
                             ("Heading 3", 11, True), ("Title", 16, True)]:
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s.font.size = Pt(size)
        s.font.bold = bold

    title = re.search(r"\\title\{(.*?)\}\s*\\titlerunning", body, re.S)
    p = doc.add_paragraph(flatten(title.group(1)), style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author = re.search(r"\\author\{(.*?)\}\s*\\authorrunning", body, re.S)
    p = doc.add_paragraph(flatten(author.group(1)))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inst = re.search(r"\\institute\{(.*?)\}\s*\\maketitle", body, re.S)
    p = doc.add_paragraph(flatten(inst.group(1)))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}",
                         body, re.S).group(1)
    kw = re.search(r"\\keywords\{(.*?)\}\s*$", abstract, re.S)
    abst = (abstract[:kw.start()] if kw else abstract)
    doc.add_paragraph("Abstract. " + flatten(abst))
    if kw:
        doc.add_paragraph("Keywords: " +
                          flatten(kw.group(1).replace("\\and", ", ")))

    # walk
    chunks = re.split(r"\n\s*\n", body[body.find("\\end{abstract}"):])
    for blk in chunks:
        blk = blk.strip()
        if (not blk or blk.startswith("\\maketitle")
                or blk.startswith("\\end{") or blk == "\\end{abstract}"):
            continue
        m = re.match(r"\\section(\*?)\{(.*?)\}\s*(.*)$", blk, re.S)
        if m and not blk.startswith("\\subsection"):
            doc.add_paragraph(flatten(m.group(2)), style="Heading 1")
            blk = re.sub(r"\\label\{[^}]*\}", "", m.group(3)).strip()
            if not blk:
                continue
        m = re.match(r"\\subsection\{(.*?)\}\s*(.*)$", blk, re.S)
        if m:
            doc.add_paragraph(flatten(m.group(1)), style="Heading 2")
            blk = re.sub(r"\\label\{[^}]*\}", "", m.group(2)).strip()
            if not blk:
                continue
        m = re.match(r"\\subsubsection\{(.*?)\}\s*(.*)$", blk, re.S)
        if m:
            para = doc.add_paragraph()
            r = para.add_run(flatten(m.group(1)) + " ")
            r.bold = True
            para.add_run(flatten(m.group(2)))
            continue
        if blk.startswith("\\begin{table}"):
            cap = re.search(r"\\caption\{(.*?)\}\s*\\label", blk, re.S)
            doc.add_paragraph(flatten(cap.group(1)), style="Heading 3")
            tm = re.search(r"\\begin\{tabular\}\{([^}]*)\}(.*?)\\end\{tabular\}",
                           blk, re.S)
            rows = []
            for line in tm.group(2).split("\\\\"):
                line = line.replace("\\hline", "").strip()
                if not line:
                    continue
                rows.append([flatten(c) for c in line.split("&")])
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j in range(ncol):
                        t.cell(i, j).text = r[j] if j < len(r) else ""
            continue
        if blk.startswith("\\begin{figure}"):
            cap = re.search(r"\\caption\{(.*?)\}\s*\\label", blk, re.S)
            img = re.search(r"\\includegraphics[^\n]*\{([^}]*)\}", blk)
            doc.add_paragraph(f"[FIGURE: {img.group(1)}]", style="Heading 3")
            doc.add_paragraph(flatten(cap.group(1)))
            continue
        if blk.startswith("\\begin{equation}"):
            tex = re.sub(r"\\label\{[^}]*\}", "", blk)
            tex = tex.replace("\\begin{equation}", "").replace("\\end{equation}", "")
            p = doc.add_paragraph(flatten(tex))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if blk.startswith("\\bibliograph"):
            doc.add_paragraph("References", style="Heading 1")
            continue
        doc.add_paragraph(flatten(blk))

    out = os.path.join(ROOT, "build", "manuscript_editable.docx")
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
